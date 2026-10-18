import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('영양학정리.xlsx', data_only=True)
target_sheets = [
    '아픈아이 종합처방', '비만', '알러지', '결석', '위장관질환',
    '만성신부전', '췌장염', '간질환', '피부병', '암', '내분비질환', '심장병'
]

for sname in target_sheets:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n=======================================================")
        print(f"SHEET: {sname}")
        print(f"=======================================================")
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if any(x is not None for x in row):
                vals = [str(x).strip() if x is not None else '' for x in row]
                while vals and vals[-1] == '':
                    vals.pop()
                print(f"L{idx:02d}: " + " | ".join(vals))
