import openpyxl
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def analyze():
    print("=== Analyzing Prescription Diet Data ===")
    filename = '처방사료데이터 정리260821.xlsx'
    if os.path.exists(filename):
        wb = openpyxl.load_workbook(filename, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            print(f"\nSheet: {sname} (Total rows: {len(rows)})")
            if not rows:
                continue
            headers = rows[0]
            valid_rows = [r for r in rows[1:] if any(x is not None for x in r)]
            print(f"Valid data rows: {len(valid_rows)}")
            for r in valid_rows[:2]:
                print("Sample:", {str(h): v for h, v in zip(headers, r) if h and v is not None})
    
    print("\n=== Analyzing Ingredients JSON ===")
    if os.path.exists('ingredients.json'):
        with open('ingredients.json', 'r', encoding='utf-8') as f:
            data = json.load(f)
        print(f"Total ingredients: {len(data)}")
        categories = {}
        for item in data:
            cat = item.get('category', '기타')
            categories[cat] = categories.get(cat, 0) + 1
        print("Categories distribution:", categories)

if __name__ == '__main__':
    analyze()
