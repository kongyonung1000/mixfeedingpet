import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_nutrition_excel():
    fname = '영양학정리.xlsx'
    wb = openpyxl.load_workbook(fname, data_only=True)
    print(f"=== Sheets in {fname} ===")
    print(wb.sheetnames)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        print(f"\n================ Sheet: {sname} (Total rows: {len(rows)}) ================")
        valid_rows = [r for r in rows if any(x is not None for x in r)]
        print(f"Non-empty rows: {len(valid_rows)}")
        for idx, r in enumerate(valid_rows[:30]):
            # Filter out trailing Nones
            trimmed = list(r)
            while trimmed and trimmed[-1] is None:
                trimmed.pop()
            print(f"Row {idx+1}: {trimmed}")

if __name__ == '__main__':
    inspect_nutrition_excel()
