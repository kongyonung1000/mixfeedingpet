import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

fname = '영양학정리.xlsx'
wb = openpyxl.load_workbook(fname, data_only=True)
print("ALL SHEETS:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=======================================================")
    print(f"SHEET: {sname}")
    print(f"=======================================================")
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if any(x is not None for x in row):
            vals = [str(x).strip() if x is not None else '' for x in row]
            while vals and vals[-1] == '':
                vals.pop()
            print(f"L{idx:02d}: " + " | ".join(vals))
